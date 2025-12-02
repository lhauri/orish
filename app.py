"""Orish Flask Application

How to run locally:
1. (Optional) Create & activate a virtual environment.
2. Install dependencies: pip install -r requirements.txt (or `pip install flask` if requirements are unavailable).
3. Create a .env file (see .env.example) and initialize the SQLite database: python init_db.py.
4. Start the development server: python app.py.
"""

import csv
import json
import os
import random
import re
import sqlite3
from collections import Counter
from datetime import datetime
from functools import wraps
from io import BytesIO

from dotenv import load_dotenv
from flask import (
    Flask,
    abort,
    flash,
    g,
    has_request_context,
    jsonify,
    redirect,
    render_template,
    request,
    session,
    url_for,
)
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

from docx import Document
from PyPDF2 import PdfReader
import openpyxl
import requests

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
DATABASE_PATH = os.path.join(BASE_DIR, "orish.db")

load_dotenv(os.path.join(BASE_DIR, ".env"))

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("ORISH_SECRET", "dev-secret-key"),
    DATABASE=DATABASE_PATH,
)

DEEPSEEK_API_KEY = os.getenv("DEEPSEEK_API_KEY")
DEEPSEEK_BASE_URL = os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com")
DEEPSEEK_MODEL = os.getenv("DEEPSEEK_MODEL", "deepseek-chat")
DEEPSEEK_TIMEOUT = 30

ALLOWED_UPLOADS = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".csv"}


def evaluate_text_answer(prompt, reference, student_answer):
    """Use DeepSeek (or a fallback) to judge free-text answers."""
    student_answer = (student_answer or "").strip()
    reference = (reference or "").strip()
    if not student_answer:
        return {
            "is_correct": False,
            "feedback": "No answer submitted.",
            "explanation": "Please provide a response so we can review it.",
        }

    fallback_correct = student_answer.lower() == reference.lower()
    base_feedback = {
        "is_correct": fallback_correct,
        "feedback": (
            "Looks good! Keep it up." if fallback_correct else f"Expected: {reference}"
        ),
        "explanation": "",
    }

    try:
        response = _deepseek_chat(
            [
                {
                    "role": "system",
                    "content": (
                        "You are an English teacher. Strictly reply with a JSON object "
                        'like {"is_correct": bool, "feedback": "...", "explanation": "..."}'
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        f"Question: {prompt}\nExpected answer: {reference}\n"
                        f"Student answer: {student_answer}\nJudge correctness for an exam."
                    ),
                },
            ],
            temperature=0.2,
        )
        text = _extract_chat_text(response)
        if not text:
            raise RuntimeError("Empty AI feedback.")
        data = json.loads(_sanitize_ai_text_payload(text))
        return {
            "is_correct": bool(data.get("is_correct")),
            "feedback": data.get("feedback") or base_feedback["feedback"],
            "explanation": data.get("explanation", ""),
        }
    except Exception as exc:  # pragma: no cover - defensive
        app.logger.warning("DeepSeek grading failed: %s", exc)
        return base_feedback


def summarize_attempt_for_teacher(exam_title, answers):
    """Ask DeepSeek for a concise teacher-facing summary."""
    if not DEEPSEEK_API_KEY or not answers:
        return None
    serialized = "\n".join(
        f"Q: {a['question']['prompt']} | Student: {a.get('selected')} | "
        f"Correct: {a['question']['correct_answer']} | Result: {a['is_correct']} | "
        f"Feedback: {a.get('feedback') or ''}"
        for a in answers
    )
    try:
        response = _deepseek_chat(
            [
                {
                    "role": "system",
                    "content": (
                        "You are a concise English teacher assistant. Summarize performance "
                        "for another teacher in <=4 sentences."
                    ),
                },
                {
                    "role": "user",
                    "content": f"Exam: {exam_title}\nDetails:\n{serialized}",
                },
            ],
            temperature=0.3,
        )
        text = _extract_chat_text(response)
        return text.strip() if text else None
    except Exception as exc:  # pragma: no cover
        app.logger.warning("DeepSeek summary failed: %s", exc)
    return None


QUESTION_SCHEMAS = {
    "vocabulary": {
        "description": "Return JSON array of objects with keys word, correct_answer, wrong1, wrong2, wrong3.",
        "columns": ["word", "correct_answer", "wrong1", "wrong2", "wrong3"],
    },
    "grammar": {
        "description": (
            "Return JSON array of objects with keys sentence_with_placeholder "
            "(use __ for blank), correct_answer, wrong1, wrong2, wrong3."
        ),
        "columns": [
            "sentence_with_placeholder",
            "correct_answer",
            "wrong1",
            "wrong2",
            "wrong3",
        ],
    },
    "translation": {
        "description": "Return JSON array of objects with keys prompt and reference_answer.",
        "columns": ["prompt", "reference_answer"],
    },
}

JSON_BLOCK_RE = re.compile(r"(\{.*\}|\[.*\])", re.DOTALL)
API_VERSION_RE = re.compile(r"/v\d+$")

FALLBACK_GENERATED_QUESTIONS = {
    "vocabulary": [
        {
            "word": "serene",
            "correct_answer": "Calm and peaceful",
            "wrong1": "Full of energy",
            "wrong2": "Extremely loud",
            "wrong3": "Difficult to find",
        },
        {
            "word": "anticipate",
            "correct_answer": "Expect or look forward to",
            "wrong1": "Forget completely",
            "wrong2": "Argue loudly",
            "wrong3": "Hide from others",
        },
        {
            "word": "versatile",
            "correct_answer": "Able to do many things well",
            "wrong1": "Afraid of change",
            "wrong2": "Hard to see",
            "wrong3": "Very expensive",
        },
    ],
    "grammar": [
        {
            "sentence_with_placeholder": "The students ___ their essays before class.",
            "correct_answer": "had finished",
            "wrong1": "finishing",
            "wrong2": "was finish",
            "wrong3": "has finished",
        },
        {
            "sentence_with_placeholder": "If she ___ earlier, we would have caught the train.",
            "correct_answer": "had left",
            "wrong1": "lefts",
            "wrong2": "has leaving",
            "wrong3": "leaves",
        },
    ],
    "translation": [
        {
            "prompt": "Translate into English: \"Ich lerne jeden Tag neue Wörter.\"",
            "reference_answer": "I learn new words every day.",
        },
        {
            "prompt": "Translate into English: \"Wir treffen uns morgen im Park.\"",
            "reference_answer": "We are meeting in the park tomorrow.",
        },
    ],
}

FALLBACK_EXAM_TEMPLATES = [
    {
        "title": "Balanced Skills Check",
        "description": "Quick assessment drawn from the built-in bank.",
        "category": "vocabulary",
        "questions": 5,
        "items": [
            {
                "prompt": "Choose the best meaning for \"resilient\".",
                "answer_type": "mcq",
                "correct_answer": "Able to recover quickly",
                "wrong1": "Afraid of speaking",
                "wrong2": "Expensive to buy",
                "wrong3": "Easy to forget",
            },
            {
                "prompt": "Select the synonym of \"ambitious\".",
                "answer_type": "mcq",
                "correct_answer": "Driven",
                "wrong1": "Careless",
                "wrong2": "Sleepy",
                "wrong3": "Salty",
            },
        ],
    },
    {
        "title": "Grammar Tune-Up",
        "description": "Targeted practice for tenses and connectors.",
        "category": "grammar",
        "questions": 5,
        "items": [
            {
                "prompt": "If it ___ tomorrow, we will stay home.",
                "answer_type": "mcq",
                "correct_answer": "rains",
                "wrong1": "rained",
                "wrong2": "rain",
                "wrong3": "was raining",
            },
            {
                "prompt": "By the time she arrived, we ___ dinner.",
                "answer_type": "mcq",
                "correct_answer": "had started",
                "wrong1": "start",
                "wrong2": "were starting",
                "wrong3": "starting",
            },
        ],
    },
]

LEGAL_SECTIONS = [
    {
        "id": "terms",
        "title": "Terms of Service",
        "eyebrow": "Terms",
        "intro": (
            "Orish is designed for learners and teachers who want to sharpen English skills. "
            "By using the platform you agree to the following guidelines."
        ),
        "clauses": [
            {
                "title": "Permitted use",
                "body": (
                    "Use Orish only for your personal or school learning workflows. "
                    "No automated scraping, mass exports, or sharing of private classroom data without consent."
                ),
            },
            {
                "title": "Account safety",
                "body": (
                    "You are responsible for safeguarding your account credentials. "
                    "Contact us immediately if you suspect unauthorized activity."
                ),
            },
            {
                "title": "Teacher dashboards",
                "body": (
                    "Admin tools may only be used with school approval. "
                    "Exports are limited to assessing the learners you actively teach."
                ),
            },
            {
                "title": "Availability",
                "body": (
                    "We strive for high uptime but may schedule maintenance or updates. "
                    "Major changes will be communicated to registered users via email."
                ),
            },
        ],
    },
    {
        "id": "privacy",
        "title": "Privacy Policy",
        "eyebrow": "Privacy",
        "intro": (
            "We collect only the data needed to deliver the learning experience and stay compliant with GDPR and Swiss law."
        ),
        "clauses": [
            {
                "title": "Data we collect",
                "body": (
                    "Profile basics (name, username, school email), quiz activity (answers, scores, feedback), "
                    "and optional uploads for AI analysis. We never store payment information."
                ),
            },
            {
                "title": "Purpose",
                "body": (
                    "To provide practice modules, tailored feedback, teacher analytics, and core security logging."
                ),
            },
            {
                "title": "Retention",
                "body": (
                    "Accounts remain active until you request deletion. "
                    "Upon request we erase data within 30 days. Exam logs are anonymized after 18 months."
                ),
            },
            {
                "title": "Processors",
                "body": (
                    "Hosting is located in the EU. AI features can use DeepSeek/OpenAI and only send necessary snippets. "
                    "We do not share data with ad networks."
                ),
            },
            {
                "title": "Your rights",
                "body": (
                    "You may request access, correction, deletion, restriction, or data portability at privacy@orish.app."
                ),
            },
        ],
    },
    {
        "id": "contact",
        "title": "Contact & Responsible Entity",
        "eyebrow": "Contact",
        "intro": (
            "Need help with legal, privacy, or account topics? Reach out and we usually respond within two business days."
        ),
        "clauses": [
            {
                "title": "Responsible organization",
                "body": "Orish Learning Collective • Lorrainestrasse 5B • 3013 Bern • Switzerland",
            },
            {
                "title": "Email",
                "body": "privacy@orish.app",
            },
        ],
    },
]


def _fallback_questions_for_category(category, max_items=3):
    templates = FALLBACK_GENERATED_QUESTIONS.get(category, [])
    if not templates:
        return []
    sample_count = min(len(templates), max_items)
    sample = random.sample(templates, sample_count)
    return [dict(item) for item in sample]


def _local_text_analysis(snippet, custom_prompt=None):
    text = (snippet or "").strip()
    if not text:
        return {
            "summary": "No text supplied for analysis.",
            "vocabulary": "",
            "grammar": "",
            "action_points": custom_prompt or "Upload a document to receive feedback.",
        }
    tokens = re.findall(r"[A-Za-z']+", text.lower())
    word_count = len(tokens)
    unique_words = len(set(tokens))
    sentences = [s.strip() for s in re.split(r"[.!?]+", text) if s.strip()]
    sentence_count = len(sentences) or 1
    first_idea = sentences[0][:160] if sentences else text[:160]
    common_words = [
        word for word, _ in Counter(tokens).most_common(3) if len(word) > 3
    ]
    summary = (
        f"Local analyzer reviewed about {word_count} words across {sentence_count} sentences. "
        f"Opening idea: {first_idea}"
    )
    vocab = (
        f"Frequently used terms: {', '.join(common_words)}"
        if common_words
        else "Vocabulary is varied; keep highlighting precise verbs."
    )
    grammar = (
        "Mix short and long sentences for better rhythm."
        if sentence_count > 3
        else "Consider adding more supporting sentences for clarity."
    )
    action = custom_prompt or "Underline confusing areas and rewrite one sentence for clarity."
    return {
        "summary": summary,
        "vocabulary": vocab,
        "grammar": grammar,
        "action_points": action,
    }


def _sanitize_ai_text_payload(content):
    """Strip markdown fences or stray whitespace before JSON parsing."""
    text = (content or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?", "", text, flags=re.IGNORECASE).strip()
        text = re.sub(r"```$", "", text).strip()
    return text


def _inform_ai_fallback(message):
    if has_request_context():
        flash(message, "info")


def _normalized_base_url():
    base = (DEEPSEEK_BASE_URL or "").strip() or "https://api.deepseek.com"
    base = base.rstrip("/")
    if not API_VERSION_RE.search(base):
        base = f"{base}/v1"
    return base


def _deepseek_chat(messages, temperature=0.4):
    if not DEEPSEEK_API_KEY:
        raise RuntimeError("AI key missing")
    url = f"{_normalized_base_url()}/chat/completions"
    payload = {
        "model": DEEPSEEK_MODEL,
        "messages": messages,
        "temperature": temperature,
    }
    headers = {
        "Authorization": f"Bearer {DEEPSEEK_API_KEY}",
        "Content-Type": "application/json",
    }
    try:
        response = requests.post(
            url,
            headers=headers,
            json=payload,
            timeout=DEEPSEEK_TIMEOUT,
        )
        response.raise_for_status()
        return response.json()
    except requests.RequestException as exc:  # pragma: no cover - network / API issues
        app.logger.warning("DeepSeek chat request failed: %s", exc)
        raise RuntimeError("AI request failed") from exc


def _extract_chat_text(response):
    if not response:
        return ""
    choices = response.get("choices") if isinstance(response, dict) else getattr(response, "choices", None)
    if not choices:
        return ""
    first_choice = choices[0]
    message = first_choice.get("message") if isinstance(first_choice, dict) else getattr(first_choice, "message", None)
    content = message.get("content", "") if isinstance(message, dict) else getattr(message, "content", message) if message else ""
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        parts = []
        for chunk in content:
            if isinstance(chunk, dict) and chunk.get("type") == "text":
                parts.append(chunk.get("text", ""))
            elif isinstance(chunk, str):
                parts.append(chunk)
        return " ".join(parts).strip()
    return str(content).strip()


def request_ai_json(system_prompt, user_prompt):
    response = _deepseek_chat(
        [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        temperature=0.4,
    )
    text = _sanitize_ai_text_payload(_extract_chat_text(response))
    if not text:
        raise RuntimeError("AI response was empty.")
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = JSON_BLOCK_RE.search(text)
        if match:
            try:
                return json.loads(match.group(1))
            except json.JSONDecodeError:
                pass
        raise RuntimeError("AI returned invalid JSON.")


def generate_questions_with_prompt(category, prompt):
    if category not in QUESTION_SCHEMAS:
        raise ValueError("Unsupported category")
    schema = QUESTION_SCHEMAS[category]
    instructions = (
        "You are helping teachers prepare English exams. "
        "Always return valid JSON and no other text. "
        f"{schema['description']} Produce 1-3 fresh questions."
    )
    user_prompt = (
        f"Category: {category}\nTeacher guidance: {prompt or 'Create standard practice.'}"
    )
    use_fallback = False
    try:
        data = request_ai_json(instructions, user_prompt)
    except RuntimeError as exc:
        app.logger.warning("AI question generation failed: %s", exc)
        data = _fallback_questions_for_category(category)
        use_fallback = True
    except Exception as exc:  # pragma: no cover
        app.logger.warning("AI generation error: %s", exc)
        data = _fallback_questions_for_category(category)
        use_fallback = True
    if isinstance(data, dict):
        data = [data]
    filtered = []
    for item in data:
        if not isinstance(item, dict):
            continue
        filtered.append({key: item.get(key, "").strip() for key in schema["columns"]})
    if not filtered:
        raise RuntimeError("AI did not return usable content.")
    if use_fallback:
        _inform_ai_fallback("AI temporarily offline. Added sample questions instead.")
    return filtered


def generate_exam_from_prompt(prompt):
    instructions = (
        "Create a single exam descriptor as JSON with keys title, description, "
        "category (vocabulary/grammar/translation), questions (int between 3 and 10) "
        "and items (array of questions). Each item needs prompt, answer_type ('mcq' or 'text'), "
        "correct_answer, wrong1, wrong2, wrong3, reference_answer. "
        "Return JSON only."
    )
    use_fallback = False
    try:
        data = request_ai_json(instructions, prompt or "Create a balanced assessment.")
    except RuntimeError as exc:
        app.logger.warning("AI exam generation failed: %s", exc)
        use_fallback = True
        data = random.choice(FALLBACK_EXAM_TEMPLATES)
    except Exception as exc:  # pragma: no cover
        app.logger.warning("Exam AI error: %s", exc)
        use_fallback = True
        data = random.choice(FALLBACK_EXAM_TEMPLATES)
    if isinstance(data, list):
        data = data[0]
    category = data.get("category", "vocabulary").lower()
    if category not in CATEGORIES:
        category = "vocabulary"
    raw_questions = data.get("questions", 5)
    try:
        questions = int(raw_questions)
    except (TypeError, ValueError):
        questions = 5
    payload = {
        "title": data.get("title", "AI Exam Draft").strip()[:80],
        "description": data.get("description", "").strip()[:200],
        "category": category,
        "questions": max(3, min(questions, 10)),
    }
    items = data.get("items") or data.get("questions") or []
    normalized = []
    for item in items:
        if not isinstance(item, dict):
            continue
        prompt_text = item.get("prompt", "").strip()
        if not prompt_text:
            continue
        answer_type = (item.get("answer_type") or "").lower()
        if answer_type not in {"mcq", "text"}:
            answer_type = "text" if category == "translation" else "mcq"
        entry = {
            "prompt": prompt_text[:400],
            "answer_type": answer_type,
            "correct_answer": (item.get("correct_answer") or "").strip()[:200],
            "wrong1": (item.get("wrong1") or "").strip()[:200],
            "wrong2": (item.get("wrong2") or "").strip()[:200],
            "wrong3": (item.get("wrong3") or "").strip()[:200],
            "reference_answer": (
                item.get("reference_answer")
                or item.get("correct_answer")
                or ""
            ).strip()[:300],
        }
        normalized.append(entry)
    payload["items"] = normalized
    if use_fallback:
        desc = payload["description"] or "Quick mixed drill."
        if prompt:
            desc += f" (Based on: {prompt[:60]})"
            payload["description"] = desc[:200]
        _inform_ai_fallback("AI unavailable. Created a built-in exam template.")
    return payload


def analyze_text_with_ai(text, custom_prompt=None):
    snippet = text[:4000]
    instructions = (
        "Provide a JSON object with keys summary, vocabulary, grammar, action_points. "
        "Each value should be short strings or bullet-like sentences."
    )
    user_content = (
        f"Student material:\n{snippet}\n\nFocus: {custom_prompt or 'Highlight strengths and improvements.'}"
    )
    try:
        data = request_ai_json(instructions, user_content)
    except RuntimeError as exc:
        app.logger.warning("AI analyzer unavailable: %s", exc)
        _inform_ai_fallback("AI analyzer offline. Showing heuristic feedback instead.")
        return _local_text_analysis(snippet, custom_prompt)
    except Exception as exc:  # pragma: no cover
        app.logger.warning("AI analysis error: %s", exc)
        _inform_ai_fallback("AI analyzer offline. Showing heuristic feedback instead.")
        return _local_text_analysis(snippet, custom_prompt)
    return {
        "summary": data.get("summary", "No summary produced."),
        "vocabulary": data.get("vocabulary", ""),
        "grammar": data.get("grammar", ""),
        "action_points": data.get("action_points", ""),
    }


def extract_text_from_upload(file_storage):
    if not file_storage or not file_storage.filename:
        raise ValueError("Please choose a file to upload.")
    filename = secure_filename(file_storage.filename)
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_UPLOADS:
        raise ValueError("Unsupported file type.")
    data = file_storage.read()
    if not data:
        raise ValueError("File appears to be empty.")
    stream = BytesIO(data)
    if ext in {".txt", ".md"}:
        return data.decode("utf-8", errors="ignore")
    if ext == ".pdf":
        reader = PdfReader(stream)
        parts = []
        for page in reader.pages[:10]:
            parts.append(page.extract_text() or "")
        return "\n".join(parts)
    if ext == ".docx":
        doc = Document(stream)
        return "\n".join(p.text for p in doc.paragraphs)
    if ext == ".xlsx":
        wb = openpyxl.load_workbook(stream, data_only=True)
        sheet = wb.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                rows.append(" ".join(cells))
        return "\n".join(rows)
    if ext == ".csv":
        stream.seek(0)
        text = data.decode("utf-8", errors="ignore")
        reader = csv.reader(text.splitlines())
        return "\n".join(" ".join(row) for row in reader)
    raise ValueError("Unsupported file type.")

CATEGORIES = {
    "vocabulary": {
        "label": "Vocabulary",
        "icon": "type",
        "table": "questions_vocabulary",
        "prompt_builder": lambda row: f"Select the correct meaning for the word '{row['word']}'.",
        "answer_type": "mcq",
    },
    "grammar": {
        "label": "Grammar",
        "icon": "book",
        "table": "questions_grammar",
        "prompt_builder": lambda row: row[
            "sentence_with_placeholder"
        ].replace("__", "____"),
        "answer_type": "mcq",
    },
    "translation": {
        "label": "Translation",
        "icon": "languages",
        "table": "questions_translation",
        "prompt_builder": lambda row: row["prompt"],
        "answer_type": "text",
    },
}


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(app.config["DATABASE"])
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def _ensure_column(db, table, column, definition):
    """Add a column if it does not exist yet."""
    existing = {
        row["name"]
        for row in db.execute(f"PRAGMA table_info({table})").fetchall()
    }
    if column not in existing:
        db.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def init_tables():
    db = get_db()
    cursor = db.cursor()
    cursor.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            is_admin INTEGER DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS questions_vocabulary (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            word TEXT NOT NULL,
            correct_answer TEXT NOT NULL,
            wrong1 TEXT NOT NULL,
            wrong2 TEXT NOT NULL,
            wrong3 TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS questions_grammar (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sentence_with_placeholder TEXT NOT NULL,
            correct_answer TEXT NOT NULL,
            wrong1 TEXT NOT NULL,
            wrong2 TEXT NOT NULL,
            wrong3 TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS questions_translation (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            prompt TEXT NOT NULL,
            reference_answer TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            category TEXT NOT NULL,
            score INTEGER NOT NULL,
            total INTEGER NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY (user_id) REFERENCES users (id)
        );

        CREATE TABLE IF NOT EXISTS exams (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            description TEXT,
            category TEXT NOT NULL,
            questions INTEGER DEFAULT 5,
            is_active INTEGER DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS exam_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            exam_id INTEGER NOT NULL,
            score INTEGER NOT NULL,
            total INTEGER NOT NULL,
            details TEXT NOT NULL,
            ai_feedback TEXT,
             mode TEXT NOT NULL DEFAULT 'test',
            created_at TEXT NOT NULL,
            FOREIGN KEY (user_id) REFERENCES users (id),
            FOREIGN KEY (exam_id) REFERENCES exams (id)
        );

        CREATE TABLE IF NOT EXISTS question_groups (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            subject TEXT NOT NULL,
            description TEXT,
            ai_prompt TEXT,
            created_by INTEGER,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (created_by) REFERENCES users (id)
        );

        CREATE TABLE IF NOT EXISTS question_group_memberships (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            group_id INTEGER NOT NULL,
            category TEXT NOT NULL,
            question_id INTEGER NOT NULL,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE (group_id, category, question_id),
            FOREIGN KEY (group_id) REFERENCES question_groups (id)
        );

        CREATE TABLE IF NOT EXISTS question_group_assignments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            group_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            can_view INTEGER DEFAULT 1,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            UNIQUE (group_id, user_id),
            FOREIGN KEY (group_id) REFERENCES question_groups (id),
            FOREIGN KEY (user_id) REFERENCES users (id)
        );

        CREATE TABLE IF NOT EXISTS exam_questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            exam_id INTEGER NOT NULL,
            prompt TEXT NOT NULL,
            answer_type TEXT NOT NULL DEFAULT 'mcq',
            correct_answer TEXT,
            wrong1 TEXT,
            wrong2 TEXT,
            wrong3 TEXT,
            reference_answer TEXT,
            position INTEGER DEFAULT 0,
            ai_source TEXT,
            FOREIGN KEY (exam_id) REFERENCES exams (id)
        );

        CREATE TABLE IF NOT EXISTS exam_assignments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            exam_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            can_study INTEGER DEFAULT 1,
            can_test INTEGER DEFAULT 1,
            UNIQUE (exam_id, user_id),
            FOREIGN KEY (exam_id) REFERENCES exams (id),
            FOREIGN KEY (user_id) REFERENCES users (id)
        );
        """
    )
    _ensure_column(db, "exams", "study_enabled", "INTEGER DEFAULT 1")
    _ensure_column(db, "exams", "test_enabled", "INTEGER DEFAULT 1")
    _ensure_column(db, "exams", "ai_prompt", "TEXT")
    _ensure_column(db, "exam_attempts", "mode", "TEXT DEFAULT 'test'")
    db.commit()


def row_value(row, key, default=None):
    """Safe helper for sqlite Row objects."""
    try:
        if key in row.keys():
            return row[key]
    except Exception:
        pass
    return default


def format_question_row(category_key, row):
    """Normalize DB rows into quiz/exam friendly payloads."""
    category = CATEGORIES[category_key]
    answer_type = category.get("answer_type", "mcq")
    prompt = category["prompt_builder"](row)
    correct = row_value(row, "correct_answer") or row_value(
        row, "reference_answer"
    )
    question = {
        "id": row["id"],
        "prompt": prompt,
        "correct_answer": correct,
        "answer_type": answer_type,
        "meta": {"source": "bank"},
    }
    if answer_type == "mcq":
        options = [
            row_value(row, "correct_answer"),
            row_value(row, "wrong1"),
            row_value(row, "wrong2"),
            row_value(row, "wrong3"),
        ]
        question["options"] = [opt for opt in options if opt]
        random.shuffle(question["options"])
    else:
        question["meta"]["reference_hint"] = row_value(row, "reference_answer")
    if category_key == "vocabulary":
        question["meta"]["word"] = row["word"]
    elif category_key == "grammar":
        question["meta"]["sentence"] = row[
            "sentence_with_placeholder"
        ].replace("__", "____")
    return question


def fetch_questions_by_ids(category_key, question_ids):
    if not question_ids:
        return []
    db = get_db()
    table = CATEGORIES[category_key]["table"]
    placeholders = ",".join(["?"] * len(question_ids))
    rows = db.execute(
        f"SELECT * FROM {table} WHERE id IN ({placeholders})",
        question_ids,
    ).fetchall()
    mapped = {row["id"]: row for row in rows}
    ordered = []
    for question_id in question_ids:
        row = mapped.get(question_id)
        if row:
            ordered.append(format_question_row(category_key, row))
    return ordered


def login_required(view):
    @wraps(view)
    def wrapped_view(**kwargs):
        if g.user is None:
            flash("Please log in to access this page.", "warning")
            return redirect(url_for("login"))
        return view(**kwargs)

    return wrapped_view


def admin_required(view):
    @wraps(view)
    def wrapped_view(**kwargs):
        if g.user is None or not g.user["is_admin"]:
            flash("Admin access required.", "danger")
            return redirect(url_for("dashboard"))
        return view(**kwargs)

    return wrapped_view


@app.before_request
def load_logged_in_user():
    user_id = session.get("user_id")
    if user_id is None:
        g.user = None
    else:
        user = (
            get_db()
            .execute("SELECT * FROM users WHERE id = ?", (user_id,))
            .fetchone()
        )
        g.user = user


@app.context_processor
def inject_globals():
    return {"current_year": datetime.utcnow().year}


def fetch_random_questions(category_key, limit=5):
    category = CATEGORIES[category_key]
    db = get_db()
    rows = db.execute(
        f"SELECT * FROM {category['table']} ORDER BY RANDOM() LIMIT ?",
        (limit,),
    ).fetchall()
    if not rows:
        raise ValueError("No questions available for this category yet. Please ask your teacher to add some.")
    return [format_question_row(category_key, row) for row in rows]


def fetch_exam_specific_question_rows(exam_id):
    db = get_db()
    return db.execute(
        """
        SELECT * FROM exam_questions
        WHERE exam_id = ?
        ORDER BY position ASC, id ASC
        """,
        (exam_id,),
    ).fetchall()


def format_exam_specific_question(row):
    answer_type = row["answer_type"] or "mcq"
    prompt = row["prompt"]
    correct_answer = (
        row_value(row, "correct_answer") or row_value(row, "reference_answer")
    )
    question = {
        "id": f"exam-{row['id']}",
        "prompt": prompt,
        "correct_answer": correct_answer,
        "answer_type": answer_type,
        "meta": {"source": "exam"},
    }
    if answer_type == "mcq":
        options = [
            row_value(row, "correct_answer"),
            row_value(row, "wrong1"),
            row_value(row, "wrong2"),
            row_value(row, "wrong3"),
        ]
        question["options"] = [opt for opt in options if opt]
        random.shuffle(question["options"])
    else:
        question["meta"]["reference_hint"] = row_value(row, "reference_answer")
    return question


def build_exam_question_set(exam_row):
    specific_rows = fetch_exam_specific_question_rows(exam_row["id"])
    specific_questions = [format_exam_specific_question(row) for row in specific_rows]
    needed = max(0, exam_row["questions"] - len(specific_questions))
    if needed:
        try:
            general_questions = fetch_random_questions(exam_row["category"], limit=needed)
        except ValueError as exc:
            if specific_questions:
                raise ValueError(
                    f"This exam only has {len(specific_questions)} custom question(s). "
                    "Add more exam-specific questions or restock the bank to continue."
                ) from exc
            raise
        specific_questions.extend(general_questions)
    if len(specific_questions) < exam_row["questions"]:
        raise ValueError(
            f"This exam needs {exam_row['questions']} questions but only {len(specific_questions)} "
            "are available. Please add more exam-specific questions or replenish the question bank."
        )
    return specific_questions


def count_general_questions():
    db = get_db()
    totals = {}
    for key, meta in CATEGORIES.items():
        totals[key] = (
            db.execute(f"SELECT COUNT(*) FROM {meta['table']}").fetchone()[0]
        )
    return totals


def exam_has_assignments(exam_id):
    db = get_db()
    row = db.execute(
        "SELECT COUNT(*) AS assigned FROM exam_assignments WHERE exam_id = ?",
        (exam_id,),
    ).fetchone()
    return bool(row and row["assigned"])


def get_exam_assignment(exam_id, user_id):
    if not user_id:
        return None
    db = get_db()
    return db.execute(
        """
        SELECT can_study, can_test
        FROM exam_assignments
        WHERE exam_id = ? AND user_id = ?
        """,
        (exam_id, user_id),
    ).fetchone()


def user_can_take_exam(exam_row, user, mode):
    if not user:
        return False
    if user["is_admin"]:
        return True
    if not exam_has_assignments(exam_row["id"]):
        return True
    assignment = get_exam_assignment(exam_row["id"], user["id"])
    if not assignment:
        return False
    if mode == "study" and not exam_row["study_enabled"]:
        return False
    if mode == "test" and not exam_row["test_enabled"]:
        return False
    can_flag = assignment["can_study"] if mode == "study" else assignment["can_test"]
    return bool(can_flag)


def load_question_group(group_id):
    return (
        get_db()
        .execute("SELECT * FROM question_groups WHERE id = ?", (group_id,))
        .fetchone()
    )


def fetch_group_questions(group_id):
    db = get_db()
    memberships = db.execute(
        """
        SELECT category, question_id
        FROM question_group_memberships
        WHERE group_id = ?
        ORDER BY id ASC
        """,
        (group_id,),
    ).fetchall()
    if not memberships:
        return []
    grouped = {}
    for row in memberships:
        category = row["category"]
        if category not in CATEGORIES:
            continue
        grouped.setdefault(category, []).append(row["question_id"])
    questions = []
    for category_key, ids in grouped.items():
        questions.extend(fetch_questions_by_ids(category_key, ids))
    return questions


def user_can_view_group(group_row, user):
    if not user:
        return False
    if user["is_admin"]:
        return True
    db = get_db()
    row = db.execute(
        """
        SELECT can_view
        FROM question_group_assignments
        WHERE group_id = ? AND user_id = ?
        """,
        (group_row["id"], user["id"]),
    ).fetchone()
    return bool(row and row["can_view"])


@app.route("/")
def home():
    return render_template("home.html")


@app.route("/legal")
def legal():
    return render_template(
        "legal.html",
        sections=LEGAL_SECTIONS,
        legal_updated=datetime.utcnow().strftime("%d.%m.%Y"),
    )


@app.route("/mindmap")
def mindmap():
    return render_template("mindmap.html")


@app.route("/analyze", methods=["GET", "POST"])
@login_required
def analyze():
    analysis = None
    extracted = ""
    custom_prompt = ""
    if request.method == "POST":
        custom_prompt = request.form.get("prompt", "").strip()
        file = request.files.get("document")
        try:
            extracted = extract_text_from_upload(file)
            analysis = analyze_text_with_ai(extracted, custom_prompt)
        except ValueError as exc:
            flash(str(exc), "warning")
        except RuntimeError as exc:
            flash(str(exc), "danger")
    return render_template(
        "analyze.html",
        analysis=analysis,
        custom_prompt=custom_prompt,
        sample_text=extracted[:1200],
    )


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm", "")

        if not username or not email or not password:
            flash("All fields are required.", "warning")
        elif password != confirm:
            flash("Passwords do not match.", "warning")
        else:
            db = get_db()
            try:
                db.execute(
                    "INSERT INTO users (username, email, password_hash) VALUES (?, ?, ?)",
                    (username, email, generate_password_hash(password)),
                )
                db.commit()
                flash("Account created! Please log in.", "success")
                return redirect(url_for("login"))
            except sqlite3.IntegrityError:
                flash("Username or email already exists.", "danger")

    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        identifier = (
            request.form.get("identifier")
            or request.form.get("email", "")
        ).strip()
        password = request.form.get("password", "")

        user = (
            get_db()
            .execute(
                "SELECT * FROM users WHERE email = ? OR username = ?",
                (identifier, identifier),
            )
            .fetchone()
        )
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            flash(f"Welcome back, {user['username']}!", "success")
            return redirect(url_for("dashboard"))
        flash("Invalid credentials.", "danger")

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("Logged out successfully.", "info")
    return redirect(url_for("home"))


@app.route("/dashboard")
@login_required
def dashboard():
    db = get_db()
    total_quizzes = db.execute(
        "SELECT COUNT(*) FROM results WHERE user_id = ?",
        (g.user["id"],),
    ).fetchone()[0]
    best_scores = {}
    for key in CATEGORIES:
        best = db.execute(
            "SELECT MAX(score) FROM results WHERE user_id = ? AND category = ?",
            (g.user["id"], key),
        ).fetchone()[0]
        best_scores[key] = best or 0
    recent_results = db.execute(
        "SELECT * FROM results WHERE user_id = ? ORDER BY created_at DESC LIMIT 5",
        (g.user["id"],),
    ).fetchall()

    return render_template(
        "dashboard.html",
        total_quizzes=total_quizzes,
        best_scores=best_scores,
        categories=CATEGORIES,
        results=recent_results,
    )


@app.route("/admin/users", methods=["GET", "POST"])
@admin_required
def admin_users():
    db = get_db()
    if request.method == "POST":
        action = request.form.get("action", "promote")
        is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
        if action == "create":
            username = request.form.get("username", "").strip()
            email = request.form.get("email", "").strip()
            password = request.form.get("password", "")
            role = request.form.get("role", "student")
            if not username or not email or not password:
                flash("All fields are required to create an account.", "warning")
                return redirect(url_for("admin_users"))
            if len(password) < 8:
                flash("Password must be at least 8 characters.", "warning")
                return redirect(url_for("admin_users"))
            is_admin = 1 if role == "teacher" else 0
            try:
                db.execute(
                    "INSERT INTO users (username, email, password_hash, is_admin) VALUES (?, ?, ?, ?)",
                    (username, email, generate_password_hash(password), is_admin),
                )
                db.commit()
                flash(f"Created {'teacher' if is_admin else 'student'} account for {username}.", "success")
            except sqlite3.IntegrityError:
                flash("Username or email already exists.", "danger")
            return redirect(url_for("admin_users"))

        try:
            target_id = int(request.form.get("user_id"))
        except (TypeError, ValueError):
            message = "Invalid user selection."
            if is_ajax:
                return jsonify({"status": "error", "message": message}), 400
            flash(message, "danger")
            return redirect(url_for("admin_users"))
        target = (
            db.execute("SELECT * FROM users WHERE id = ?", (target_id,)).fetchone()
        )
        if not target:
            message = "User not found."
            if is_ajax:
                return jsonify({"status": "error", "message": message}), 404
            flash(message, "warning")
            return redirect(url_for("admin_users"))

        def ajax_response(message, status="ok", extra=None, status_code=200):
            payload = {"status": status, "message": message, "user_id": target_id}
            if extra:
                payload.update(extra)
            response = jsonify(payload)
            response.status_code = status_code
            return response

        if action == "promote":
            if target["is_admin"]:
                message = f"{target['username']} is already a teacher."
                if is_ajax:
                    return ajax_response(message, status="info")
                flash(message, "info")
            else:
                db.execute(
                    "UPDATE users SET is_admin = 1 WHERE id = ?",
                    (target_id,),
                )
                db.commit()
                message = f"{target['username']} is now a teacher."
                if is_ajax:
                    return ajax_response(message, extra={"role": "teacher"})
                flash(message, "success")
            return redirect(url_for("admin_users"))
        if action == "demote":
            if target_id == g.user["id"]:
                message = "You cannot remove your own teacher role."
                if is_ajax:
                    return ajax_response(message, status="error", status_code=400)
                flash(message, "warning")
            elif not target["is_admin"]:
                message = f"{target['username']} is already a student."
                if is_ajax:
                    return ajax_response(message, status="info")
                flash(message, "info")
            else:
                db.execute(
                    "UPDATE users SET is_admin = 0 WHERE id = ?",
                    (target_id,),
                )
                db.commit()
                message = f"{target['username']} was set to student."
                if is_ajax:
                    return ajax_response(message, extra={"role": "student"})
                flash(message, "success")
            return redirect(url_for("admin_users"))
        if action == "delete":
            if target_id == g.user["id"]:
                message = "You cannot delete your own account."
                if is_ajax:
                    return ajax_response(message, status="error", status_code=400)
                flash(message, "warning")
                return redirect(url_for("admin_users"))
            db.execute("DELETE FROM results WHERE user_id = ?", (target_id,))
            db.execute("DELETE FROM exam_attempts WHERE user_id = ?", (target_id,))
            db.execute("DELETE FROM users WHERE id = ?", (target_id,))
            db.commit()
            message = f"Deleted {target['username']} and their data."
            if is_ajax:
                return ajax_response(message, extra={"role": "teacher" if target["is_admin"] else "student"})
            flash(message, "success")
            return redirect(url_for("admin_users"))
        if action == "prepare_delete":
            message = f"Confirm delete {target['username']}."
            if is_ajax:
                return ajax_response(message, status="info")
            return message, 200
        else:
            message = "Unknown action."
            if is_ajax:
                return ajax_response(message, status="error", status_code=400)
            flash(message, "warning")
            return redirect(url_for("admin_users"))

    users = db.execute(
        "SELECT id, username, email, is_admin FROM users ORDER BY username ASC"
    ).fetchall()
    teacher_rows = [user for user in users if user["is_admin"]]
    student_rows = [user for user in users if not user["is_admin"]]
    return render_template(
        "admin_users.html",
        teachers=teacher_rows,
        students=student_rows,
    )


@app.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    db = get_db()
    if request.method == "POST":
        action = request.form.get("profile_action", "password")
        if action == "username":
            new_username = request.form.get("new_username", "").strip()
            if not new_username:
                flash("Please enter a username.", "warning")
            elif len(new_username) < 3:
                flash("Username must be at least 3 characters.", "warning")
            else:
                existing = db.execute(
                    "SELECT id FROM users WHERE username = ? AND id != ?",
                    (new_username, g.user["id"]),
                ).fetchone()
                if existing:
                    flash("That username is already taken.", "danger")
                else:
                    db.execute(
                        "UPDATE users SET username = ? WHERE id = ?",
                        (new_username, g.user["id"]),
                    )
                    db.commit()
                    g.user = (
                        db.execute("SELECT * FROM users WHERE id = ?", (g.user["id"],))
                        .fetchone()
                    )
                    flash("Username updated.", "success")
            return redirect(url_for("profile"))
        else:
            current_password = request.form.get("current_password", "")
            new_password = request.form.get("new_password", "")
            confirm_password = request.form.get("confirm_password", "")
            if not current_password or not new_password or not confirm_password:
                flash("Please complete all password fields.", "warning")
            elif not check_password_hash(g.user["password_hash"], current_password):
                flash("Current password is incorrect.", "danger")
            elif len(new_password) < 8:
                flash("New password must be at least 8 characters.", "warning")
            elif new_password != confirm_password:
                flash("New passwords do not match.", "warning")
            else:
                new_hash = generate_password_hash(new_password)
                db.execute(
                    "UPDATE users SET password_hash = ? WHERE id = ?",
                    (new_hash, g.user["id"]),
                )
                db.commit()
                g.user = (
                    db.execute("SELECT * FROM users WHERE id = ?", (g.user["id"],))
                    .fetchone()
                )
                flash("Password updated successfully.", "success")
            return redirect(url_for("profile"))

    results = db.execute(
        "SELECT * FROM results WHERE user_id = ? ORDER BY created_at DESC LIMIT 10",
        (g.user["id"],),
    ).fetchall()
    exam_attempts = db.execute(
        """
        SELECT ea.*, e.title
        FROM exam_attempts ea
        JOIN exams e ON e.id = ea.exam_id
        WHERE ea.user_id = ?
        ORDER BY ea.created_at DESC
        LIMIT 5
        """,
        (g.user["id"],),
    ).fetchall()

    stats_row = db.execute(
        """
        SELECT
            COUNT(*) AS attempts,
            COALESCE(SUM(score), 0) AS total_score,
            COALESCE(SUM(total), 0) AS total_possible
        FROM results
        WHERE user_id = ?
        """,
        (g.user["id"],),
    ).fetchone()
    exam_total = (
        db.execute(
            "SELECT COUNT(*) FROM exam_attempts WHERE user_id = ?",
            (g.user["id"],),
        ).fetchone()[0]
    )
    accuracy = 0
    if stats_row["total_possible"]:
        accuracy = round(
            (stats_row["total_score"] / stats_row["total_possible"]) * 100, 1
        )
    last_activity = None
    if results:
        last_activity = results[0]["created_at"]
    elif exam_attempts:
        last_activity = exam_attempts[0]["created_at"]

    category_rows = db.execute(
        """
        SELECT
            category,
            COUNT(*) AS attempts,
            AVG(CASE WHEN total > 0 THEN CAST(score AS REAL) / total END) * 100 AS accuracy
        FROM results
        WHERE user_id = ?
        GROUP BY category
        """,
        (g.user["id"],),
    ).fetchall()
    category_stats = [
        {
            "category": row["category"],
            "label": CATEGORIES.get(row["category"], {}).get("label", row["category"]),
            "attempts": row["attempts"],
            "accuracy": round(row["accuracy"] or 0, 1) if row["accuracy"] else 0,
        }
        for row in category_rows
    ]
    profile_stats = {
        "attempts": stats_row["attempts"] or 0,
        "accuracy": accuracy,
        "exam_attempts": exam_total,
        "last_activity": last_activity,
    }
    role_label = "Teacher" if g.user["is_admin"] else "Student"
    return render_template(
        "profile.html",
        results=results,
        exam_attempts=exam_attempts,
        profile_stats=profile_stats,
        category_stats=category_stats,
        role_label=role_label,
    )


@app.route("/quiz/select")
@login_required
def quiz_select():
    return render_template("quiz_select.html", categories=CATEGORIES)


@app.route("/exams")
@login_required
def exams():
    db = get_db()
    general_counts = count_general_questions()
    if g.user["is_admin"]:
        exam_rows = db.execute(
            "SELECT e.*, 1 AS can_study, 1 AS can_test FROM exams e ORDER BY e.id DESC"
        ).fetchall()
    else:
        exam_rows = db.execute(
            """
            SELECT e.*, ea.can_study, ea.can_test
            FROM exams e
            LEFT JOIN exam_assignments ea
                ON ea.exam_id = e.id AND ea.user_id = ?
            WHERE e.is_active = 1
            ORDER BY e.id DESC
            """,
            (g.user["id"],),
        ).fetchall()
    exams = []
    for row in exam_rows:
        data = dict(row)
        category_meta = CATEGORIES.get(data["category"], {})
        table = category_meta.get("table")
        available_general = general_counts.get(data["category"], 0) if table else 0
        specific_total = (
            db.execute(
                "SELECT COUNT(*) FROM exam_questions WHERE exam_id = ?", (data["id"],)
            ).fetchone()[0]
        )
        assigned_count = db.execute(
            "SELECT COUNT(*) FROM exam_assignments WHERE exam_id = ?",
            (data["id"],),
        ).fetchone()[0]
        available = available_general + specific_total
        data["category_label"] = category_meta.get("label", data["category"].title())
        data["category_icon"] = category_meta.get("icon", "book")
        data["has_questions"] = available >= data["questions"]
        data["available_questions"] = available
        data["general_questions"] = available_general
        data["specific_questions"] = specific_total
        data["assigned_count"] = assigned_count
        data["missing_questions"] = max(0, data["questions"] - available)
        if g.user["is_admin"]:
            data["can_study"] = True
            data["can_test"] = True
        else:
            assignment_exists = row["can_study"] is not None or row["can_test"] is not None
            open_to_all = assigned_count == 0
            if not assignment_exists and not open_to_all:
                # Teachers restricted this exam; skip if not assigned.
                continue
            data["can_study"] = (
                bool(row["can_study"]) if assignment_exists else open_to_all
            )
            data["can_test"] = (
                bool(row["can_test"]) if assignment_exists else open_to_all
            )
        exams.append(data)
    attempts = db.execute(
        """
        SELECT ea.*, e.title
        FROM exam_attempts ea
        JOIN exams e ON e.id = ea.exam_id
        WHERE ea.user_id = ?
        ORDER BY ea.created_at DESC
        LIMIT 10
        """,
        (g.user["id"],),
    ).fetchall()
    return render_template("exams.html", exams=exams, attempts=attempts, categories=CATEGORIES)


@app.route("/exams/create", methods=["POST"])
@admin_required
def create_exam():
    title = request.form.get("title", "").strip()
    description = request.form.get("description", "").strip()
    category = request.form.get("category", "vocabulary")
    try:
        questions = int(request.form.get("questions", 5))
    except ValueError:
        questions = 5
    study_enabled = 1 if request.form.get("study_enabled", "on") else 0
    test_enabled = 1 if request.form.get("test_enabled", "on") else 0
    if not title or category not in CATEGORIES:
        flash("Please provide a valid title and category.", "warning")
        return redirect(url_for("exams"))
    questions = max(3, min(questions, 15))
    db = get_db()
    db.execute(
        """
        INSERT INTO exams (title, description, category, questions, study_enabled, test_enabled)
        VALUES (?, ?, ?, ?, ?, ?)
        """,
        (title, description, category, questions, study_enabled, test_enabled),
    )
    db.commit()
    flash("Exam created. Add exam-specific questions and share it with students.", "success")
    return redirect(url_for("exams"))


@app.route("/admin/exams/generate", methods=["POST"])
@admin_required
def generate_exam_ai():
    prompt = request.form.get("prompt", "").strip()
    try:
        payload = generate_exam_from_prompt(prompt)
    except RuntimeError as exc:
        flash(str(exc), "danger")
        return redirect(url_for("exams"))
    db = get_db()
    category = payload["category"]
    questions = max(3, min(payload["questions"], 15))
    cur = db.execute(
        """
        INSERT INTO exams (title, description, category, questions, study_enabled, test_enabled, ai_prompt)
        VALUES (?, ?, ?, ?, 1, 1, ?)
        """,
        (
            payload["title"],
            payload["description"],
            category,
            questions,
            prompt or None,
        ),
    )
    exam_id = cur.lastrowid
    ai_items = payload.get("items") or []
    selected_items = ai_items[:questions]
    for position, item in enumerate(selected_items, start=1):
        answer_type = "text" if item.get("answer_type") == "text" else "mcq"
        db.execute(
            """
            INSERT INTO exam_questions (exam_id, prompt, answer_type, correct_answer, wrong1, wrong2, wrong3, reference_answer, position, ai_source)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                exam_id,
                item.get("prompt") or "",
                answer_type,
                item.get("correct_answer"),
                item.get("wrong1"),
                item.get("wrong2"),
                item.get("wrong3"),
                item.get("reference_answer"),
                position,
                "ai",
            ),
        )
    db.commit()
    flash(
        f"AI created exam '{payload['title']}' with {len(selected_items)} custom question(s).",
        "success",
    )
    return redirect(url_for("exams"))


def load_exam(exam_id):
    exam = (
        get_db()
        .execute("SELECT * FROM exams WHERE id = ?", (exam_id,))
        .fetchone()
    )
    return exam


@app.route("/exams/<int:exam_id>/manage")
@admin_required
def manage_exam(exam_id):
    exam = load_exam(exam_id)
    if not exam:
        flash("Exam not found.", "warning")
        return redirect(url_for("exams"))
    db = get_db()
    question_rows = fetch_exam_specific_question_rows(exam_id)
    questions = [
        {
            "id": row["id"],
            "prompt": row["prompt"],
            "answer_type": row["answer_type"],
            "correct_answer": row["correct_answer"],
            "wrong1": row["wrong1"],
            "wrong2": row["wrong2"],
            "wrong3": row["wrong3"],
            "reference_answer": row["reference_answer"],
            "ai_source": row["ai_source"],
            "position": row["position"],
        }
        for row in question_rows
    ]
    assignments = db.execute(
        """
        SELECT ea.*, u.username, u.email
        FROM exam_assignments ea
        JOIN users u ON u.id = ea.user_id
        WHERE ea.exam_id = ?
        ORDER BY u.username ASC
        """,
        (exam_id,),
    ).fetchall()
    stats = {
        "specific": len(questions),
        "general": count_general_questions().get(exam["category"], 0),
    }
    return render_template(
        "exam_manage.html",
        exam=exam,
        questions=questions,
        assignments=assignments,
        stats=stats,
        categories=CATEGORIES,
    )


@app.route("/exams/<int:exam_id>/settings", methods=["POST"])
@admin_required
def update_exam_settings(exam_id):
    exam = load_exam(exam_id)
    if not exam:
        flash("Exam not found.", "warning")
        return redirect(url_for("exams"))
    try:
        questions = int(request.form.get("questions", exam["questions"]))
    except (TypeError, ValueError):
        questions = exam["questions"]
    questions = max(3, min(questions, 30))
    study_enabled = 1 if request.form.get("study_enabled") else 0
    test_enabled = 1 if request.form.get("test_enabled") else 0
    is_active = 1 if request.form.get("is_active") else 0
    db = get_db()
    db.execute(
        """
        UPDATE exams
        SET questions = ?, study_enabled = ?, test_enabled = ?, is_active = ?
        WHERE id = ?
        """,
        (questions, study_enabled, test_enabled, is_active, exam_id),
    )
    db.commit()
    flash("Exam settings updated.", "success")
    return redirect(url_for("manage_exam", exam_id=exam_id))


def _next_exam_question_position(exam_id):
    db = get_db()
    row = db.execute(
        "SELECT COALESCE(MAX(position), 0) + 1 AS next_pos FROM exam_questions WHERE exam_id = ?",
        (exam_id,),
    ).fetchone()
    return row["next_pos"] if row else 1


@app.route("/exams/<int:exam_id>/questions", methods=["POST"])
@admin_required
def add_exam_question(exam_id):
    exam = load_exam(exam_id)
    if not exam:
        flash("Exam not found.", "warning")
        return redirect(url_for("exams"))
    prompt = request.form.get("prompt", "").strip()
    answer_type = request.form.get("answer_type", "mcq")
    if answer_type not in {"mcq", "text"}:
        answer_type = "mcq"
    db = get_db()
    position = _next_exam_question_position(exam_id)
    if not prompt:
        flash("Please provide a prompt.", "warning")
        return redirect(url_for("manage_exam", exam_id=exam_id))
    if answer_type == "text":
        reference = request.form.get("reference_answer", "").strip()
        if not reference:
            flash("Text questions need a reference answer.", "warning")
            return redirect(url_for("manage_exam", exam_id=exam_id))
        correct_answer = request.form.get("correct_answer", "").strip() or reference
        wrongs = (None, None, None)
    else:
        correct_answer = request.form.get("correct_answer", "").strip()
        wrongs = [request.form.get(f"wrong{i}", "").strip() for i in range(1, 4)]
        if not correct_answer or not all(wrongs):
            flash("Multiple-choice questions need one correct and three incorrect options.", "warning")
            return redirect(url_for("manage_exam", exam_id=exam_id))
        reference = ""
    db.execute(
        """
        INSERT INTO exam_questions
        (exam_id, prompt, answer_type, correct_answer, wrong1, wrong2, wrong3, reference_answer, position, ai_source)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            exam_id,
            prompt,
            answer_type,
            correct_answer,
            wrongs[0] if answer_type == "mcq" else None,
            wrongs[1] if answer_type == "mcq" else None,
            wrongs[2] if answer_type == "mcq" else None,
            reference,
            position,
            "manual",
        ),
    )
    db.commit()
    flash("Added exam-specific question.", "success")
    return redirect(url_for("manage_exam", exam_id=exam_id))


@app.route("/exams/<int:exam_id>/questions/ai", methods=["POST"])
@admin_required
def add_exam_questions_ai(exam_id):
    exam = load_exam(exam_id)
    if not exam:
        flash("Exam not found.", "warning")
        return redirect(url_for("exams"))
    prompt = request.form.get("prompt", "").strip()
    try:
        generated = generate_questions_with_prompt(exam["category"], prompt)
    except RuntimeError as exc:
        flash(str(exc), "danger")
        return redirect(url_for("manage_exam", exam_id=exam_id))
    db = get_db()
    position = _next_exam_question_position(exam_id)
    inserted = 0
    for item in generated:
        if exam["category"] == "translation":
            answer_type = "text"
            question_prompt = item.get("prompt")
            reference_answer = item.get("reference_answer")
            correct_answer = reference_answer
            wrongs = (None, None, None)
        elif exam["category"] == "grammar":
            answer_type = "mcq"
            question_prompt = item.get("sentence_with_placeholder", "Complete the sentence.")
            correct_answer = item.get("correct_answer")
            wrongs = (item.get("wrong1"), item.get("wrong2"), item.get("wrong3"))
            reference_answer = ""
        else:
            answer_type = "mcq"
            word = item.get("word", "this word")
            question_prompt = f"What is the best meaning of \"{word}\"?"
            correct_answer = item.get("correct_answer")
            wrongs = (item.get("wrong1"), item.get("wrong2"), item.get("wrong3"))
            reference_answer = ""
        if not question_prompt or not correct_answer:
            continue
        db.execute(
            """
            INSERT INTO exam_questions
            (exam_id, prompt, answer_type, correct_answer, wrong1, wrong2, wrong3, reference_answer, position, ai_source)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                exam_id,
                question_prompt,
                answer_type,
                correct_answer,
                wrongs[0],
                wrongs[1],
                wrongs[2],
                reference_answer,
                position,
                "ai",
            ),
        )
        position += 1
        inserted += 1
    db.commit()
    flash(f"Added {inserted} AI question(s) to this exam.", "success")
    return redirect(url_for("manage_exam", exam_id=exam_id))


@app.route("/exams/<int:exam_id>/questions/<int:question_id>/delete", methods=["POST"])
@admin_required
def delete_exam_question(exam_id, question_id):
    exam = load_exam(exam_id)
    if not exam:
        flash("Exam not found.", "warning")
        return redirect(url_for("exams"))
    db = get_db()
    db.execute(
        "DELETE FROM exam_questions WHERE id = ? AND exam_id = ?",
        (question_id, exam_id),
    )
    db.commit()
    flash("Removed exam question.", "info")
    return redirect(url_for("manage_exam", exam_id=exam_id))


@app.route("/exams/<int:exam_id>/assign", methods=["POST"])
@admin_required
def assign_exam_to_student(exam_id):
    exam = load_exam(exam_id)
    if not exam:
        flash("Exam not found.", "warning")
        return redirect(url_for("exams"))
    identifier = request.form.get("identifier", "").strip()
    if not identifier:
        flash("Enter a student username or email.", "warning")
        return redirect(url_for("manage_exam", exam_id=exam_id))
    db = get_db()
    user = db.execute(
        "SELECT * FROM users WHERE username = ? OR email = ?",
        (identifier, identifier),
    ).fetchone()
    if not user:
        flash("No matching user found.", "warning")
        return redirect(url_for("manage_exam", exam_id=exam_id))
    can_study = 1 if request.form.get("can_study") else 0
    can_test = 1 if request.form.get("can_test") else 0
    existing = db.execute(
        "SELECT id FROM exam_assignments WHERE exam_id = ? AND user_id = ?",
        (exam_id, user["id"]),
    ).fetchone()
    if not can_study and not can_test:
        db.execute(
            "DELETE FROM exam_assignments WHERE exam_id = ? AND user_id = ?",
            (exam_id, user["id"]),
        )
        db.commit()
        flash(f"Removed {user['username']} from this exam.", "info")
        return redirect(url_for("manage_exam", exam_id=exam_id))
    if existing:
        db.execute(
            "UPDATE exam_assignments SET can_study = ?, can_test = ? WHERE id = ?",
            (can_study, can_test, existing["id"]),
        )
    else:
        db.execute(
            "INSERT INTO exam_assignments (exam_id, user_id, can_study, can_test) VALUES (?, ?, ?, ?)",
            (exam_id, user["id"], can_study, can_test),
        )
    db.commit()
    flash(f"Shared exam with {user['username']}.", "success")
    return redirect(url_for("manage_exam", exam_id=exam_id))


@app.route("/exams/<int:exam_id>/assign/<int:assignment_id>/delete", methods=["POST"])
@admin_required
def delete_exam_assignment(exam_id, assignment_id):
    exam = load_exam(exam_id)
    if not exam:
        flash("Exam not found.", "warning")
        return redirect(url_for("exams"))
    db = get_db()
    db.execute(
        "DELETE FROM exam_assignments WHERE id = ? AND exam_id = ?",
        (assignment_id, exam_id),
    )
    db.commit()
    flash("Removed assignment.", "info")
    return redirect(url_for("manage_exam", exam_id=exam_id))


@app.route("/exams/<int:exam_id>/take", methods=["GET", "POST"])
@login_required
def take_exam(exam_id):
    exam = load_exam(exam_id)
    if not exam or not exam["is_active"]:
        flash("This exam is no longer available.", "warning")
        return redirect(url_for("exams"))
    mode = request.args.get("mode", request.form.get("mode", "test") or "test").lower()
    if mode not in {"study", "test"}:
        mode = "test"
    if mode == "study" and not exam["study_enabled"]:
        flash("Study mode is disabled for this exam.", "warning")
        return redirect(url_for("exams"))
    if mode == "test" and not exam["test_enabled"]:
        flash("This exam is not accepting test attempts right now.", "warning")
        return redirect(url_for("exams"))
    if not user_can_take_exam(exam, g.user, mode):
        flash("This exam is not shared with you for that mode.", "warning")
        return redirect(url_for("exams"))
    exam_state = session.get("exam")
    if (
        not exam_state
        or exam_state.get("exam_id") != exam_id
        or exam_state.get("mode") != mode
    ):
        try:
            start_exam_session(exam, mode)
        except ValueError as exc:
            flash(str(exc), "warning")
            return redirect(url_for("exams"))
        exam_state = session["exam"]

    if request.method == "POST":
        exam_state = session.get("exam")
        current_index = exam_state["current"]
        question = exam_state["questions"][current_index]
        if question["answer_type"] == "text":
            selected = request.form.get("text_answer", "").strip()
            evaluation = evaluate_text_answer(
                question["prompt"], question["correct_answer"], selected
            )
            is_correct = evaluation["is_correct"]
            feedback = evaluation.get("feedback")
            explanation = evaluation.get("explanation")
        else:
            selected = request.form.get("answer")
            if not selected:
                flash("Please pick an option to continue.", "warning")
                return redirect(url_for("take_exam", exam_id=exam_id, mode=mode))
            is_correct = selected == question["correct_answer"]
            feedback = ""
            explanation = ""
        exam_state["answers"].append(
            {
                "question": question,
                "selected": selected,
                "is_correct": is_correct,
                "feedback": feedback,
                "explanation": explanation,
            }
        )
        if is_correct:
            exam_state["score"] += 1
        exam_state["current"] += 1
        session["exam"] = exam_state

        if exam_state["current"] >= exam_state["total"]:
            db = get_db()
            details_json = json.dumps(exam_state["answers"])
            ai_summary = None
            if mode == "test":
                ai_summary = summarize_attempt_for_teacher(
                    exam_state["title"], exam_state["answers"]
                )
            cursor = db.execute(
                """
                INSERT INTO exam_attempts (user_id, exam_id, score, total, details, ai_feedback, mode, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    g.user["id"],
                    exam_id,
                    exam_state["score"],
                    exam_state["total"],
                    details_json,
                    ai_summary,
                    mode,
                    datetime.utcnow().isoformat(),
                ),
            )
            db.commit()
            attempt_id = cursor.lastrowid
            session.pop("exam", None)
            return redirect(url_for("exam_result", attempt_id=attempt_id))

    return render_template(
        "exam_take.html",
        exam=exam,
        mode=mode,
        exam_state=exam_state,
        question=exam_state["questions"][exam_state["current"]],
        current=exam_state["current"] + 1,
        total=exam_state["total"],
        categories=CATEGORIES,
    )


@app.route("/exams/results/<int:attempt_id>")
@login_required
def exam_result(attempt_id):
    db = get_db()
    attempt = db.execute(
        """
        SELECT ea.*, e.title, e.category, u.username
        FROM exam_attempts ea
        JOIN exams e ON e.id = ea.exam_id
        JOIN users u ON u.id = ea.user_id
        WHERE ea.id = ?
        """,
        (attempt_id,),
    ).fetchone()
    if not attempt:
        flash("Exam attempt not found.", "warning")
        return redirect(url_for("exams"))
    if attempt["user_id"] != g.user["id"] and not g.user["is_admin"]:
        flash("You do not have access to that report.", "danger")
        return redirect(url_for("dashboard"))
    answers = json.loads(attempt["details"])
    return render_template(
        "exam_results.html",
        attempt=attempt,
        answers=answers,
        category=CATEGORIES.get(attempt["category"]),
    )


@app.route("/admin/exams/attempts")
@admin_required
def admin_exam_attempts():
    db = get_db()
    attempts = db.execute(
        """
        SELECT ea.*, e.title, u.username
        FROM exam_attempts ea
        JOIN exams e ON e.id = ea.exam_id
        JOIN users u ON u.id = ea.user_id
        ORDER BY ea.created_at DESC
        LIMIT 50
        """
    ).fetchall()
    return render_template("admin_exam_attempts.html", attempts=attempts)


@app.route("/admin/exams/attempts/<int:attempt_id>")
@admin_required
def admin_exam_attempt_detail(attempt_id):
    db = get_db()
    attempt = db.execute(
        """
        SELECT ea.*, e.title, e.category, u.username, u.email
        FROM exam_attempts ea
        JOIN exams e ON e.id = ea.exam_id
        JOIN users u ON u.id = ea.user_id
        WHERE ea.id = ?
        """,
        (attempt_id,),
    ).fetchone()
    if not attempt:
        flash("Attempt not found.", "warning")
        return redirect(url_for("admin_exam_attempts"))
    answers = json.loads(attempt["details"])
    return render_template(
        "exam_attempt_detail.html",
        attempt=attempt,
        answers=answers,
        category=CATEGORIES.get(attempt["category"]),
    )


def start_quiz_session(category_key):
    questions = fetch_random_questions(category_key, limit=5)
    session["quiz"] = {
        "category": category_key,
        "questions": questions,
        "current": 0,
        "score": 0,
        "answers": [],
        "total": len(questions),
    }


def start_group_session(group_row):
    questions = fetch_group_questions(group_row["id"])
    if not questions:
        raise ValueError("This study pack has no questions yet.")
    session["group_quiz"] = {
        "group_id": group_row["id"],
        "group_name": group_row["name"],
        "group_subject": group_row["subject"],
        "questions": questions,
        "current": 0,
        "score": 0,
        "answers": [],
        "total": len(questions),
    }


def start_exam_session(exam_row, mode):
    try:
        questions = build_exam_question_set(exam_row)
    except ValueError as exc:
        raise ValueError(str(exc))
    session["exam"] = {
        "exam_id": exam_row["id"],
        "title": exam_row["title"],
        "category": exam_row["category"],
        "questions": questions,
        "current": 0,
        "score": 0,
        "answers": [],
        "total": len(questions),
        "mode": mode,
    }


@app.route("/quiz/<category>", methods=["GET", "POST"])
@login_required
def quiz(category):
    if category not in CATEGORIES:
        flash("Unknown category.", "danger")
        return redirect(url_for("quiz_select"))

    quiz_state = session.get("quiz")
    if not quiz_state or quiz_state.get("category") != category:
        try:
            start_quiz_session(category)
        except ValueError as exc:
            flash(str(exc), "warning")
            return redirect(url_for("quiz_select"))
        quiz_state = session["quiz"]

    if request.method == "POST":
        quiz_state = session.get("quiz")
        current_index = quiz_state["current"]
        question = quiz_state["questions"][current_index]
        if question["answer_type"] == "text":
            selected = request.form.get("text_answer", "").strip()
            evaluation = evaluate_text_answer(
                question["prompt"], question["correct_answer"], selected
            )
            is_correct = evaluation["is_correct"]
            feedback = evaluation.get("feedback")
            explanation = evaluation.get("explanation")
        else:
            selected = request.form.get("answer")
            if not selected:
                flash("Please pick an option to continue.", "warning")
                return redirect(url_for("quiz", category=category))
            is_correct = selected == question["correct_answer"]
            feedback = ""
            explanation = ""
        quiz_state["answers"].append(
            {
                "question": question,
                "selected": selected,
                "is_correct": is_correct,
                "feedback": feedback,
                "explanation": explanation,
            }
        )
        if is_correct:
            quiz_state["score"] += 1
        quiz_state["current"] += 1
        session["quiz"] = quiz_state

        if quiz_state["current"] >= quiz_state["total"]:
            db = get_db()
            db.execute(
                "INSERT INTO results (user_id, category, score, total, created_at) VALUES (?, ?, ?, ?, ?)",
                (
                    g.user["id"],
                    category,
                    quiz_state["score"],
                    quiz_state["total"],
                    datetime.utcnow().isoformat(),
                ),
            )
            db.commit()
            session["quiz_result"] = quiz_state
            session.pop("quiz", None)
            return redirect(url_for("results"))

    quiz_state = session.get("quiz")
    current_index = quiz_state["current"]
    question = quiz_state["questions"][current_index]

    return render_template(
        "quiz.html",
        category=category,
        category_meta=CATEGORIES[category],
        question=question,
        current=current_index + 1,
        total=quiz_state["total"],
    )


@app.route("/study-packs")
@login_required
def study_packs():
    db = get_db()
    if g.user["is_admin"]:
        groups = db.execute(
            """
            SELECT g.*, COUNT(DISTINCT m.id) AS question_count,
                   COUNT(DISTINCT qa.id) AS student_count
            FROM question_groups g
            LEFT JOIN question_group_memberships m ON m.group_id = g.id
            LEFT JOIN question_group_assignments qa ON qa.group_id = g.id
            ORDER BY g.created_at DESC
            """
        ).fetchall()
    else:
        groups = db.execute(
            """
            SELECT g.*, COUNT(DISTINCT m.id) AS question_count
            FROM question_group_assignments qa
            JOIN question_groups g ON g.id = qa.group_id
            LEFT JOIN question_group_memberships m ON m.group_id = g.id
            WHERE qa.user_id = ? AND qa.can_view = 1
            GROUP BY g.id
            ORDER BY g.created_at DESC
            """,
            (g.user["id"],),
        ).fetchall()
    return render_template("study_packs.html", groups=groups, categories=CATEGORIES)


@app.route("/study-packs/<int:group_id>", methods=["GET", "POST"])
@login_required
def study_group(group_id):
    group = load_question_group(group_id)
    if not group:
        flash("Study pack not found.", "warning")
        return redirect(url_for("study_packs"))
    if not user_can_view_group(group, g.user):
        flash("You do not have access to that study pack.", "danger")
        return redirect(url_for("study_packs"))
    pack_state = session.get("group_quiz")
    if not pack_state or pack_state.get("group_id") != group_id:
        try:
            start_group_session(group)
        except ValueError as exc:
            flash(str(exc), "warning")
            return redirect(url_for("study_packs"))
        pack_state = session["group_quiz"]
    if request.method == "POST":
        current_index = pack_state["current"]
        question = pack_state["questions"][current_index]
        if question["answer_type"] == "text":
            selected = request.form.get("text_answer", "").strip()
            evaluation = evaluate_text_answer(
                question["prompt"], question["correct_answer"], selected
            )
            is_correct = evaluation["is_correct"]
            feedback = evaluation.get("feedback")
            explanation = evaluation.get("explanation")
        else:
            selected = request.form.get("answer")
            if not selected:
                flash("Please pick an option to continue.", "warning")
                return redirect(url_for("study_group", group_id=group_id))
            is_correct = selected == question["correct_answer"]
            feedback = ""
            explanation = ""
        pack_state["answers"].append(
            {
                "question": question,
                "selected": selected,
                "is_correct": is_correct,
                "feedback": feedback,
                "explanation": explanation,
            }
        )
        if is_correct:
            pack_state["score"] += 1
        pack_state["current"] += 1
        session["group_quiz"] = pack_state
        if pack_state["current"] >= pack_state["total"]:
            db = get_db()
            db.execute(
                "INSERT INTO results (user_id, category, score, total, created_at) VALUES (?, ?, ?, ?, ?)",
                (
                    g.user["id"],
                    group["subject"],
                    pack_state["score"],
                    pack_state["total"],
                    datetime.utcnow().isoformat(),
                ),
            )
            db.commit()
            result_payload = dict(pack_state)
            result_payload["category"] = group["subject"]
            result_payload["group_name"] = group["name"]
            result_payload["group_id"] = group_id
            session["quiz_result"] = result_payload
            session.pop("group_quiz", None)
            return redirect(url_for("results"))
    current_index = pack_state["current"]
    question = pack_state["questions"][current_index]
    return render_template(
        "quiz.html",
        category=group["subject"],
        category_meta=CATEGORIES[group["subject"]],
        question=question,
        current=current_index + 1,
        total=pack_state["total"],
        group=group,
    )


@app.route("/results")
@login_required
def results():
    quiz_result = session.get("quiz_result")
    if not quiz_result:
        flash("No quiz data to show.", "info")
        return redirect(url_for("dashboard"))
    return render_template("results.html", quiz=quiz_result, category=CATEGORIES[quiz_result["category"]])


@app.route("/admin/questions")
@admin_required
def admin_questions():
    category = request.args.get("category", "vocabulary")
    if category not in CATEGORIES:
        category = "vocabulary"
    db = get_db()
    rows = db.execute(
        f"SELECT * FROM {CATEGORIES[category]['table']} ORDER BY id DESC"
    ).fetchall()
    group_rows = db.execute(
        """
        SELECT g.*, 
               COUNT(DISTINCT m.id) AS question_count,
               COUNT(DISTINCT qa.id) AS student_count
        FROM question_groups g
        LEFT JOIN question_group_memberships m ON m.group_id = g.id
        LEFT JOIN question_group_assignments qa ON qa.group_id = g.id
        WHERE g.subject = ?
        GROUP BY g.id
        ORDER BY g.created_at DESC
        """,
        (category,),
    ).fetchall()
    membership_map = {}
    assignment_map = {}
    if group_rows:
        group_ids = [row["id"] for row in group_rows]
        placeholders = ",".join(["?"] * len(group_ids))
        membership_rows = db.execute(
            f"""
            SELECT m.group_id, m.question_id, g.name
            FROM question_group_memberships m
            JOIN question_groups g ON g.id = m.group_id
            WHERE m.group_id IN ({placeholders})
            """,
            group_ids,
        ).fetchall()
        for item in membership_rows:
            membership_map.setdefault(item["question_id"], []).append(
                {"name": item["name"], "group_id": item["group_id"]}
            )
        assignment_rows = db.execute(
            f"""
            SELECT qa.id, qa.group_id, u.username, u.email
            FROM question_group_assignments qa
            JOIN users u ON u.id = qa.user_id
            WHERE qa.group_id IN ({placeholders})
            ORDER BY u.username ASC
            """,
            group_ids,
        ).fetchall()
        for row in assignment_rows:
            assignment_map.setdefault(row["group_id"], []).append(row)
    return render_template(
        "admin_questions.html",
        category=category,
        categories=CATEGORIES,
        rows=rows,
        groups=group_rows,
        memberships=membership_map,
        group_assignments=assignment_map,
    )


@app.route("/admin/questions/<category>/add", methods=["POST"])
@admin_required
def add_question(category):
    if category not in CATEGORIES:
        abort(404)
    table = CATEGORIES[category]["table"]
    db = get_db()

    if category == "vocabulary":
        word = request.form.get("word", "").strip()
        correct = request.form.get("correct_answer", "").strip()
        wrongs = [request.form.get(f"wrong{i}", "").strip() for i in range(1, 4)]
        if word and correct and all(wrongs):
            db.execute(
                f"INSERT INTO {table} (word, correct_answer, wrong1, wrong2, wrong3) VALUES (?, ?, ?, ?, ?)",
                (word, correct, wrongs[0], wrongs[1], wrongs[2]),
            )
    elif category == "grammar":
        sentence = request.form.get("sentence", "").strip()
        correct = request.form.get("correct_answer", "").strip()
        wrongs = [request.form.get(f"wrong{i}", "").strip() for i in range(1, 4)]
        if sentence and correct and all(wrongs):
            db.execute(
                f"INSERT INTO {table} (sentence_with_placeholder, correct_answer, wrong1, wrong2, wrong3) VALUES (?, ?, ?, ?, ?)",
                (sentence, correct, wrongs[0], wrongs[1], wrongs[2]),
            )
    else:  # translation
        prompt = request.form.get("prompt", "").strip()
        reference = request.form.get("reference_answer", "").strip()
        if prompt and reference:
            db.execute(
                f"INSERT INTO {table} (prompt, reference_answer) VALUES (?, ?)",
                (prompt, reference),
            )
    db.commit()
    flash("Question added.", "success")
    return redirect(url_for("admin_questions", category=category))


@app.route("/admin/questions/<category>/generate", methods=["POST"])
@admin_required
def generate_question_ai(category):
    if category not in CATEGORIES:
        abort(404)
    prompt = request.form.get("prompt", "").strip()
    try:
        generated = generate_questions_with_prompt(category, prompt)
    except RuntimeError as exc:
        flash(str(exc), "danger")
        return redirect(url_for("admin_questions", category=category))
    table = CATEGORIES[category]["table"]
    db = get_db()
    for item in generated:
        if category == "vocabulary":
            db.execute(
                f"INSERT INTO {table} (word, correct_answer, wrong1, wrong2, wrong3) VALUES (?, ?, ?, ?, ?)",
                (
                    item["word"],
                    item["correct_answer"],
                    item["wrong1"],
                    item["wrong2"],
                    item["wrong3"],
                ),
            )
        elif category == "grammar":
            db.execute(
                f"INSERT INTO {table} (sentence_with_placeholder, correct_answer, wrong1, wrong2, wrong3) VALUES (?, ?, ?, ?, ?)",
                (
                    item["sentence_with_placeholder"],
                    item["correct_answer"],
                    item["wrong1"],
                    item["wrong2"],
                    item["wrong3"],
                ),
            )
        else:
            db.execute(
                f"INSERT INTO {table} (prompt, reference_answer) VALUES (?, ?)",
                (item["prompt"], item["reference_answer"]),
            )
    db.commit()
    flash(f"Generated {len(generated)} question(s).", "success")
    return redirect(url_for("admin_questions", category=category))


@app.route("/admin/questions/<category>/<int:question_id>/delete", methods=["POST"])
@admin_required
def delete_question(category, question_id):
    if category not in CATEGORIES:
        abort(404)
    table = CATEGORIES[category]["table"]
    db = get_db()
    db.execute(f"DELETE FROM {table} WHERE id = ?", (question_id,))
    db.commit()
    flash("Question deleted.", "info")
    return redirect(url_for("admin_questions", category=category))


@app.route("/admin/questions/<category>/<int:question_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_question(category, question_id):
    if category not in CATEGORIES:
        abort(404)
    table = CATEGORIES[category]["table"]
    db = get_db()
    question = db.execute(
        f"SELECT * FROM {table} WHERE id = ?",
        (question_id,),
    ).fetchone()
    if not question:
        flash("Question not found.", "warning")
        return redirect(url_for("admin_questions", category=category))

    if request.method == "POST":
        if category == "vocabulary":
            fields = (
                request.form.get("word", "").strip(),
                request.form.get("correct_answer", "").strip(),
                request.form.get("wrong1", "").strip(),
                request.form.get("wrong2", "").strip(),
                request.form.get("wrong3", "").strip(),
                question_id,
            )
            db.execute(
                f"UPDATE {table} SET word=?, correct_answer=?, wrong1=?, wrong2=?, wrong3=? WHERE id = ?",
                fields,
            )
        elif category == "grammar":
            fields = (
                request.form.get("sentence", "").strip(),
                request.form.get("correct_answer", "").strip(),
                request.form.get("wrong1", "").strip(),
                request.form.get("wrong2", "").strip(),
                request.form.get("wrong3", "").strip(),
                question_id,
            )
            db.execute(
                f"UPDATE {table} SET sentence_with_placeholder=?, correct_answer=?, wrong1=?, wrong2=?, wrong3=? WHERE id = ?",
                fields,
            )
        else:  # translation
            fields = (
                request.form.get("prompt", "").strip(),
                request.form.get("reference_answer", "").strip(),
                question_id,
            )
            db.execute(
                f"UPDATE {table} SET prompt=?, reference_answer=? WHERE id = ?",
                fields,
            )
        db.commit()
        flash("Question updated.", "success")
        return redirect(url_for("admin_questions", category=category))

    return render_template(
        "edit_question.html",
        category=category,
        categories=CATEGORIES,
        question=question,
    )


@app.route("/admin/question-groups/<category>/create", methods=["POST"])
@admin_required
def create_question_group(category):
    if category not in CATEGORIES:
        abort(404)
    name = request.form.get("name", "").strip()
    description = request.form.get("description", "").strip()
    ai_prompt = request.form.get("ai_prompt", "").strip()
    if not name:
        flash("Name your question group.", "warning")
        return redirect(url_for("admin_questions", category=category))
    db = get_db()
    db.execute(
        """
        INSERT INTO question_groups (name, subject, description, ai_prompt, created_by)
        VALUES (?, ?, ?, ?, ?)
        """,
        (name, category, description, ai_prompt or None, g.user["id"]),
    )
    db.commit()
    flash("Created new question group.", "success")
    return redirect(url_for("admin_questions", category=category))


@app.route("/admin/question-groups/<category>/assign-question", methods=["POST"])
@admin_required
def assign_question_to_group(category):
    if category not in CATEGORIES:
        abort(404)
    try:
        question_id = int(request.form.get("question_id", 0))
    except (TypeError, ValueError):
        question_id = 0
    if not question_id:
        flash("Select a question to assign.", "warning")
        return redirect(url_for("admin_questions", category=category))
    try:
        group_id = int(request.form.get("group_id", 0))
    except (TypeError, ValueError):
        group_id = 0
    group = load_question_group(group_id)
    if not group or group["subject"] != category:
        flash("Group not found for that subject.", "warning")
        return redirect(url_for("admin_questions", category=category))
    table = CATEGORIES[category]["table"]
    db = get_db()
    exists = db.execute(
        f"SELECT id FROM {table} WHERE id = ?",
        (question_id,),
    ).fetchone()
    if not exists:
        flash("Question not found.", "warning")
        return redirect(url_for("admin_questions", category=category))
    db.execute(
        """
        INSERT OR IGNORE INTO question_group_memberships (group_id, category, question_id)
        VALUES (?, ?, ?)
        """,
        (group_id, category, question_id),
    )
    db.commit()
    flash("Question added to the group.", "success")
    return redirect(url_for("admin_questions", category=category))


@app.route("/admin/question-groups/<category>/remove-question", methods=["POST"])
@admin_required
def remove_question_from_group(category):
    if category not in CATEGORIES:
        abort(404)
    try:
        group_id = int(request.form.get("group_id", 0))
        question_id = int(request.form.get("question_id", 0))
    except (TypeError, ValueError):
        group_id = 0
        question_id = 0
    if not group_id or not question_id:
        flash("Missing group or question selection.", "warning")
        return redirect(url_for("admin_questions", category=category))
    db = get_db()
    db.execute(
        """
        DELETE FROM question_group_memberships
        WHERE group_id = ? AND question_id = ? AND category = ?
        """,
        (group_id, question_id, category),
    )
    db.commit()
    flash("Removed question from group.", "info")
    return redirect(url_for("admin_questions", category=category))


@app.route("/admin/question-groups/<int:group_id>/share", methods=["POST"])
@admin_required
def share_question_group(group_id):
    group = load_question_group(group_id)
    category = request.form.get("category", "vocabulary")
    if not group:
        flash("Group not found.", "warning")
        return redirect(url_for("admin_questions", category=category))
    identifier = request.form.get("identifier", "").strip()
    if not identifier:
        flash("Enter a student username or email.", "warning")
        return redirect(url_for("admin_questions", category=group["subject"]))
    db = get_db()
    user = db.execute(
        "SELECT * FROM users WHERE username = ? OR email = ?",
        (identifier, identifier),
    ).fetchone()
    if not user:
        flash("No matching user found.", "warning")
        return redirect(url_for("admin_questions", category=group["subject"]))
    db.execute(
        """
        INSERT INTO question_group_assignments (group_id, user_id, can_view)
        VALUES (?, ?, 1)
        ON CONFLICT(group_id, user_id) DO UPDATE SET can_view = 1
        """,
        (group_id, user["id"]),
    )
    db.commit()
    flash(f"Shared '{group['name']}' with {user['username']}.", "success")
    return redirect(url_for("admin_questions", category=group["subject"]))


@app.route("/admin/question-groups/<int:group_id>/revoke/<int:assignment_id>", methods=["POST"])
@admin_required
def revoke_question_group_assignment(group_id, assignment_id):
    group = load_question_group(group_id)
    category = group["subject"] if group else "vocabulary"
    db = get_db()
    db.execute(
        "DELETE FROM question_group_assignments WHERE id = ? AND group_id = ?",
        (assignment_id, group_id),
    )
    db.commit()
    flash("Removed student from the group.", "info")
    return redirect(url_for("admin_questions", category=category))


@app.errorhandler(404)
def not_found(error):
    return (
        render_template("error.html", code=404, message="The page you requested was not found."),
        404,
    )


@app.errorhandler(500)
def server_error(error):  # pragma: no cover - best effort
    return (
        render_template(
            "error.html",
            code=500,
            message="Something unexpected happened. Please try again in a moment.",
        ),
        500,
    )


if __name__ == "__main__":
    with app.app_context():
        init_tables()
    app.run(debug=True)
